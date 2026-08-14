"""GSAP file import — routing / operation data for work order planning."""
from __future__ import annotations

import csv
import io
import re
import uuid
import xml.etree.ElementTree as ET
from datetime import date, datetime, timedelta
from html.parser import HTMLParser
from typing import Any, Optional

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy import or_
from sqlalchemy.orm import Session

from ..auth import get_current_user, require_role
from ..models import GsapSync, get_db, now_ist

router = APIRouter(prefix="/api/gsap-sync", tags=["gsap-sync"])

MAX_UPLOAD_BYTES = 16 * 1024 * 1024
ALLOWED_EXTENSIONS = (".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".tsv", ".htm", ".html", ".xml")
OLE_MAGIC = b"\xd0\xcf\x11\xe0"
ZIP_MAGIC = b"PK"

# SAP ALV / GSAP header labels → internal field (after _norm_header)
HEADER_MAP: dict[str, str] = {
    "material": "material",
    "material number": "material",
    "material no": "material",
    "materialnummer": "material",
    "matnr": "material",
    "part": "material",
    "part no": "material",
    "part number": "material",
    "article": "material",
    "article no": "material",
    "article number": "material",
    "plant": "plant",
    "plant code": "plant",
    "werks": "plant",
    "werk": "plant",
    "created on": "created_on",
    "created date": "created_on",
    "creation date": "created_on",
    "ersda": "created_on",
    "erdat": "created_on",
    "valid from": "valid_from",
    "valid from date": "valid_from",
    "datuv": "valid_from",
    "gueltig ab": "valid_from",
    "operation": "operation",
    "operation activity": "operation",
    "operation number": "operation",
    "operation no": "operation",
    "activity": "operation",
    "op": "operation",
    "vornr": "operation",
    "vorgang": "operation",
    "work centre": "work_centre",
    "work center": "work_centre",
    "workcentre": "work_centre",
    "workcenter": "work_centre",
    "work ctr": "work_centre",
    "arbpl": "work_centre",
    "arbeitsplatz": "work_centre",
    "op short text": "op_short_text",
    "oper short text": "op_short_text",
    "operation short text": "op_short_text",
    "operation text": "op_short_text",
    "activity text": "op_short_text",
    "ltxa1": "op_short_text",
    "vorgangstext": "op_short_text",
    "setup time": "setup_time",
    "set up time": "setup_time",
    "setup": "setup_time",
    "vgw01": "setup_time",
    "ruestzeit": "setup_time",
    "machine time": "machine_time",
    "mach time": "machine_time",
    "process time": "machine_time",
    "vgw02": "machine_time",
    "maschinenzeit": "machine_time",
}

COMPACT_HEADER_MAP = {key.replace(" ", ""): field for key, field in HEADER_MAP.items()}

SKIP_HEADERS = {
    "material description", "material desc", "material text", "maktx",
    "material short text",
}

JUNK_MATERIAL = {
    "material", "matnr", "total", "sum", "gesamt", "grand total",
    "dynamic list display",
}


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ")
    for src, dst in (("ü", "ue"), ("ö", "oe"), ("ä", "ae"), ("ß", "ss"),
                     ("Ü", "ue"), ("Ö", "oe"), ("Ä", "ae")):
        text = text.replace(src, dst)
    text = text.lower()
    text = re.sub(r"[\r\n]+", " ", text)
    text = re.sub(r"[./_\\]+", " ", text)
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _map_header(value: Any) -> Optional[str]:
    key = _norm_header(value)
    if not key or key in SKIP_HEADERS:
        return None
    if any(token in key for token in ("cost", "rate", "foh", "aoh", "per unit")):
        return None
    if key in HEADER_MAP:
        return HEADER_MAP[key]
    compact = key.replace(" ", "")
    return COMPACT_HEADER_MAP.get(compact)


def _parse_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
        except (OverflowError, ValueError):
            return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def _cell_str(value: Any) -> Optional[str]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).replace("\xa0", " ").strip() or None


def _is_junk_material(material: str) -> bool:
    key = _norm_header(material)
    if key in JUNK_MATERIAL:
        return True
    if re.fullmatch(r"[-=*_]+", material.strip()):
        return True
    if "dynamic list display" in key:
        return True
    return False


def _serialize_row(row: GsapSync) -> dict:
    return {
        "id": row.id,
        "material": row.material,
        "plant": row.plant,
        "created_on": str(row.created_on) if row.created_on else None,
        "valid_from": str(row.valid_from) if row.valid_from else None,
        "operation": row.operation,
        "work_centre": row.work_centre,
        "op_short_text": row.op_short_text,
        "setup_time": row.setup_time,
        "machine_time": row.machine_time,
        "upload_batch_id": row.upload_batch_id,
        "uploaded_at": row.uploaded_at,
    }


def _header_mapping(row: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    seen: set[str] = set()
    for c, cell in enumerate(row):
        field = _map_header(cell)
        if field and field not in seen:
            mapping[c] = field
            seen.add(field)
    return mapping


def _find_header_row(rows: list[list[Any]]) -> tuple[int, dict[int, str]]:
    best: Optional[tuple[int, dict[int, str]]] = None
    best_score = -1
    for r, row in enumerate(rows[:80]):
        mapping = _header_mapping(row)
        if "material" not in mapping.values():
            continue
        score = len(mapping)
        if score > best_score:
            best_score = score
            best = (r, mapping)
    if not best:
        raise HTTPException(
            400,
            "Could not find a SAP header row with a Material column. "
            "Export from SAP ALV as Spreadsheet, Local file (unconverted/tab), "
            "HTML, or CSV. Expected columns include Material, Plant, Operation, "
            "Work Centre, Op. Short Text, Setup Time, Machine Time.",
        )
    return best


def _rows_to_records(grid: list[list[Any]]) -> list[dict]:
    header_idx, col_map = _find_header_row(grid)
    records: list[dict] = []
    for row in grid[header_idx + 1 :]:
        mapped = _header_mapping(row)
        if "material" in mapped.values() and len(mapped) >= 3:
            continue  # repeated SAP page header
        raw: dict[str, Any] = {}
        empty = True
        for col_idx, field in col_map.items():
            val = row[col_idx] if col_idx < len(row) else None
            if val not in (None, ""):
                empty = False
            raw[field] = val
        if empty:
            continue
        material = _cell_str(raw.get("material"))
        if not material or _is_junk_material(material):
            continue
        records.append({
            "material": material,
            "plant": _cell_str(raw.get("plant")),
            "created_on": _parse_date(raw.get("created_on")),
            "valid_from": _parse_date(raw.get("valid_from")),
            "operation": _cell_str(raw.get("operation")),
            "work_centre": _cell_str(raw.get("work_centre")),
            "op_short_text": _cell_str(raw.get("op_short_text")),
            "setup_time": _cell_str(raw.get("setup_time")),
            "machine_time": _cell_str(raw.get("machine_time")),
        })
    if not records:
        raise HTTPException(400, "No data rows found after the SAP header row")
    return records


def _score_grid(grid: list[list[Any]]) -> int:
    try:
        _, col_map = _find_header_row(grid)
    except HTTPException:
        return -1
    return len(col_map)


def _pick_best_grid(grids: list[list[list[Any]]], empty_msg: str) -> list[list[Any]]:
    best = None
    best_score = -1
    for grid in grids:
        score = _score_grid(grid)
        if score > best_score:
            best_score = score
            best = grid
    if best is None or best_score < 0:
        raise HTTPException(400, empty_msg)
    return best


def _decode_text(content: bytes) -> str:
    if content.startswith(b"\xff\xfe"):
        return content.decode("utf-16-le")
    if content.startswith(b"\xfe\xff"):
        return content.decode("utf-16-be")
    if content.startswith(b"\xef\xbb\xbf"):
        return content.decode("utf-8-sig")
    # SAP sometimes writes UTF-16 without a BOM
    if len(content) >= 4 and content[1:2] == b"\x00" and content[3:4] == b"\x00":
        try:
            return content.decode("utf-16-le")
        except UnicodeDecodeError:
            pass
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return content.decode("cp1252")
        except UnicodeDecodeError:
            return content.decode("latin-1")


def _sniff_delimiter(text: str) -> str:
    sample_lines = [ln for ln in text.splitlines() if ln.strip()][:20]
    sample = "\n".join(sample_lines)
    counts = {
        "\t": sample.count("\t"),
        ",": sample.count(","),
        ";": sample.count(";"),
        "|": sample.count("|"),
    }
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else ","


def _grid_from_delimited(content: bytes) -> list[list[Any]]:
    text = _decode_text(content)
    delimiter = _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [list(row) for row in reader]


def _grid_from_xlsx(content: bytes) -> list[list[Any]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(400, f"Invalid Excel file: {exc}") from exc
    grids: list[list[list[Any]]] = []
    for sheet in wb.worksheets:
        grids.append([list(row) for row in sheet.iter_rows(values_only=True)])
    wb.close()
    if not grids:
        raise HTTPException(400, "Excel workbook has no sheets")
    return _pick_best_grid(
        grids,
        "Could not find a Material header in any worksheet of the Excel file",
    )


def _grid_from_ole_xls(content: bytes) -> list[list[Any]]:
    try:
        import xlrd  # type: ignore
    except ImportError as exc:
        raise HTTPException(
            400,
            "This looks like a binary Excel 97-2003 workbook. "
            "SAP ALV usually saves Spreadsheet / Local file as text, HTML, or tab-separated .xls, "
            "which are supported. Re-export from SAP as Spreadsheet, HTML, or CSV, "
            "or save this file as .xlsx.",
        ) from exc
    try:
        book = xlrd.open_workbook(file_contents=content)
        grids = []
        for i in range(book.nsheets):
            sheet = book.sheet_by_index(i)
            grids.append([sheet.row_values(r) for r in range(sheet.nrows)])
    except Exception as exc:
        raise HTTPException(400, f"Invalid .xls file: {exc}") from exc
    return _pick_best_grid(grids, "Could not find a Material header in the .xls workbook")


class _HtmlTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._table: Optional[list[list[str]]] = None
        self._row: Optional[list[str]] = None
        self._cell: Optional[list[str]] = None

    def handle_starttag(self, tag: str, attrs) -> None:
        tag = tag.lower()
        if tag == "table":
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(re.sub(r"\s+", " ", " ".join(self._cell)).strip())
            self._cell = None
        elif tag == "tr" and self._table is not None and self._row is not None:
            self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            self.tables.append(self._table)
            self._table = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def _extract_mhtml_html(text: str) -> str:
    parts = re.split(r"\r?\n--[^\n]*", text)
    for part in parts:
        if re.search(r"content-type:\s*text/html", part, re.I):
            body = re.split(r"\r?\n\r?\n", part, maxsplit=1)
            if len(body) == 2:
                return body[1]
    return text


def _xml_local(tag: str) -> str:
    return tag.split("}", 1)[-1].lower()


def _grid_from_html(content: bytes) -> list[list[Any]]:
    text = _decode_text(content)
    head = text.lstrip()[:80].lower()
    if head.startswith("mime-version") or "multipart/" in head:
        text = _extract_mhtml_html(text)
    parser = _HtmlTableParser()
    try:
        parser.feed(text)
        parser.close()
    except Exception as exc:
        raise HTTPException(400, f"Could not read SAP HTML export: {exc}") from exc
    if not parser.tables:
        raise HTTPException(400, "SAP HTML export has no table")
    return _pick_best_grid(
        parser.tables,
        "Could not find a Material column in the SAP HTML table",
    )


def _grid_from_xml_ss(content: bytes) -> list[list[Any]]:
    text = _decode_text(content)
    try:
        root = ET.fromstring(text)
    except ET.ParseError as exc:
        raise HTTPException(400, f"Invalid Excel XML export: {exc}") from exc
    grids: list[list[list[Any]]] = []
    for table_el in root.iter():
        if _xml_local(table_el.tag) != "table":
            continue
        grid: list[list[Any]] = []
        for row_el in table_el:
            if _xml_local(row_el.tag) != "row":
                continue
            row: list[Any] = []
            for cell in row_el:
                if _xml_local(cell.tag) != "cell":
                    continue
                idx = None
                for attr_name, attr_val in cell.attrib.items():
                    if _xml_local(attr_name) == "index":
                        try:
                            idx = int(attr_val)
                        except ValueError:
                            idx = None
                data_text = ""
                for child in cell:
                    if _xml_local(child.tag) == "data":
                        data_text = "".join(child.itertext()).strip()
                if idx:
                    while len(row) < idx - 1:
                        row.append("")
                row.append(data_text)
            grid.append(row)
        if grid:
            grids.append(grid)
    if not grids:
        raise HTTPException(400, "Excel XML export has no table")
    return _pick_best_grid(grids, "Could not find a Material column in the Excel XML export")


def _looks_like_html(text: str) -> bool:
    if _looks_like_xml_ss(text):
        return False
    head = text.lstrip()[:4000].lower()
    return (
        "<html" in head
        or "<!doctype html" in head
        or "<td" in head
        or "<th" in head
        or head.startswith("mime-version")
        or "content-type: multipart" in head
    )


def _looks_like_xml_ss(text: str) -> bool:
    head = text.lstrip()[:4000].lower()
    return "<workbook" in head and "spreadsheet" in head


def parse_gsap_file(content: bytes, filename: str = "") -> list[dict]:
    name = (filename or "").lower()
    if content.startswith(ZIP_MAGIC) or name.endswith((".xlsx", ".xlsm")):
        return _rows_to_records(_grid_from_xlsx(content))
    if content.startswith(OLE_MAGIC):
        return _rows_to_records(_grid_from_ole_xls(content))

    text = _decode_text(content)
    if _looks_like_xml_ss(text) or (name.endswith(".xml") and "spreadsheet" in text[:4000].lower()):
        return _rows_to_records(_grid_from_xml_ss(content))
    if _looks_like_html(text) or name.endswith((".htm", ".html")):
        return _rows_to_records(_grid_from_html(content))
    if name.endswith((".csv", ".xls", ".txt", ".tsv")) or content.startswith(
        (b"\xff\xfe", b"\xfe\xff", b"\xef\xbb\xbf")
    ):
        return _rows_to_records(_grid_from_delimited(content))

    errors: list[str] = []
    for loader in (_grid_from_xlsx, _grid_from_html, _grid_from_delimited):
        try:
            return _rows_to_records(loader(content))
        except HTTPException as exc:
            errors.append(str(exc.detail))
    raise HTTPException(
        400,
        "Could not read this SAP export. Try Spreadsheet, Local file (tab), HTML, or CSV. "
        + (errors[0] if errors else ""),
    )


def parse_gsap_excel(content: bytes) -> list[dict]:
    return parse_gsap_file(content)


class GsapUploadResult(BaseModel):
    batch_id: str
    imported: int
    replaced_previous: bool


@router.get("/")
def list_gsap_rows(
    search: Optional[str] = None,
    material: Optional[str] = None,
    limit: int = 2000,
    db: Session = Depends(get_db),
    _=Depends(require_role("admin", "superadmin", "supervisor")),
):
    q = db.query(GsapSync).order_by(GsapSync.material, GsapSync.operation, GsapSync.id)
    if material:
        q = q.filter(GsapSync.material == material)
    if search:
        like = f"%{search.strip()}%"
        q = q.filter(
            or_(
                GsapSync.material.like(like),
                GsapSync.op_short_text.like(like),
                GsapSync.work_centre.like(like),
            )
        )
    rows = q.limit(min(limit, 5000)).all()
    batch_id = rows[0].upload_batch_id if rows else None
    return {
        "items": [_serialize_row(r) for r in rows],
        "count": len(rows),
        "upload_batch_id": batch_id,
    }


@router.get("/materials")
def list_gsap_materials(
    search: Optional[str] = None,
    db: Session = Depends(get_db),
    _=Depends(get_current_user),
):
    q = db.query(GsapSync.material).distinct().order_by(GsapSync.material)
    if search:
        q = q.filter(GsapSync.material.like(f"%{search.strip()}%"))
    return [{"material": m[0]} for m in q.limit(500).all()]


@router.get("/{row_id}")
def get_gsap_row(
    row_id: int,
    db: Session = Depends(get_db),
    _=Depends(get_current_user),
):
    row = db.query(GsapSync).filter(GsapSync.id == row_id).first()
    if not row:
        raise HTTPException(404, "GSAP row not found")
    return _serialize_row(row)


@router.post("/upload", response_model=GsapUploadResult)
async def upload_gsap_excel(
    file: UploadFile = File(...),
    replace: bool = True,
    db: Session = Depends(get_db),
    user=Depends(require_role("admin", "superadmin", "supervisor")),
):
    name = (file.filename or "").lower()
    if name and not name.endswith(ALLOWED_EXTENSIONS):
        raise HTTPException(
            400,
            "Upload a SAP GSAP export (.xlsx, .xls, .csv, .txt, .html, or .xml)",
        )
    content = await file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(400, "File exceeds 16 MB limit")
    parsed = parse_gsap_file(content, file.filename or "")
    batch_id = str(uuid.uuid4())
    replaced = False
    if replace:
        db.query(GsapSync).delete()
        replaced = True
    now = now_ist()
    for item in parsed:
        db.add(GsapSync(
            material=item["material"],
            plant=item["plant"],
            created_on=item["created_on"],
            valid_from=item["valid_from"],
            operation=item["operation"],
            work_centre=item["work_centre"],
            op_short_text=item["op_short_text"],
            setup_time=item["setup_time"],
            machine_time=item["machine_time"],
            upload_batch_id=batch_id,
            uploaded_by=user.id,
            uploaded_at=now,
        ))
    db.commit()
    return GsapUploadResult(batch_id=batch_id, imported=len(parsed), replaced_previous=replaced)


@router.delete("/")
def clear_gsap_data(
    db: Session = Depends(get_db),
    _=Depends(require_role("admin", "superadmin", "supervisor")),
):
    count = db.query(GsapSync).count()
    db.query(GsapSync).delete()
    db.commit()
    return {"deleted": count}
